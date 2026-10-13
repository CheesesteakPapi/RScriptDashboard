"""Write immutable report runs and update a manifest pointer last."""
from __future__ import annotations

import csv
import hashlib
import json
import os
import tempfile
from pathlib import Path

from .core import DETAIL_FIELDS, METRIC_FIELDS, REJECT_FIELDS, Report, ReportError, summarize

AOR_DIMS = ("Component", "AOR")
COMPONENT_DIMS = ("Component", "Component Name")
OWNER_DIMS = ("Owner Badge ID", "Owner First Name", "Owner Last Name", "Owner Role",
              "Component", "Component Name", "AOR")


def csv_file(path: Path, fields, rows) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)


def workbook(path: Path, report: Report, tables: dict, moderate: int, critical: int,
             sample: bool) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ReportError("Excel output requires: python -m pip install -r requirements.txt") from exc
    book = Workbook()
    instructions = book.active
    instructions.title = "Instructions"
    for text in (
        "SYNTHETIC DEMO DATA" if sample else "Uncategorized evidence report",
        f"Snapshot: {report.summary['Snapshot Date']}",
        f"Normal: 0–{moderate - 1} calendar days since upload",
        f"Moderate: {moderate}–{critical - 1} calendar days since upload",
        f"Critical: {critical}+ calendar days since upload",
        "Age is not a three-business-day policy compliance calculation.",
        "Missing/future upload dates and missing evidence IDs are excluded into Data Quality.",
        "Owner summaries keep different badge/name/organization combinations separate.",
        "No uncategorized records." if not report.details else "Details are sorted oldest first.",
    ):
        instructions.append([text])
    instructions.column_dimensions["A"].width = 100
    sheet_specs = [
        ("Uncategorized", "Uncategorized", *tables["uncategorized.csv"]),
        ("AOR Summary", "AORSummary", *tables["aor_summary.csv"]),
        ("Component Summary", "ComponentSummary", *tables["component_summary.csv"]),
        ("Owner Summary", "OwnerSummary", *tables["owner_summary.csv"]),
        ("Snapshot", "Snapshot", *tables["snapshot.csv"]),
        ("Data Quality", "DataQuality", *tables["rejected_rows.csv"]),
    ]
    for name, table_name, fields, rows in sheet_specs:
        if len(rows) > 1_048_575:
            raise ReportError("Excel row limit exceeded; use CSV output for this dataset.")
        sheet = book.create_sheet(name)
        sheet.append(list(fields))
        for row in rows:
            sheet.append([row[key] for key in fields])
            # Source strings remain literal strings, including leading '='.
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243746")
        sheet.freeze_panes = "A2"
        last_column = get_column_letter(len(fields))
        sheet.auto_filter.ref = f"A1:{last_column}{max(1, sheet.max_row)}"
        # Excel tables are also the service-importable named data sources.
        table = Table(displayName=table_name, ref=sheet.auto_filter.ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
        for index, field in enumerate(fields, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(40, max(14, len(field) + 2))
        if "Age Band" in fields:
            age_column = list(fields).index("Days Since Upload") + 1
            color_column = list(fields).index("Age Band Color") + 1
            for row_index in range(2, sheet.max_row + 1):
                cell = sheet.cell(row_index, age_column)
                color = sheet.cell(row_index, color_column).value.lstrip("#")
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(color="FFFFFF" if color == "C41230" else "000000")
    book.save(path)


def publish_local(report: Report, output: Path, source_hash: str, source_date: str | None,
                  moderate: int = 15, critical: int = 30, excel: bool = False,
                  sample: bool = False) -> Path:
    run_id, day, fy = (report.summary[key] for key in ("Run ID", "Snapshot Date", "Fiscal Year"))
    relative = f"snapshots/FY{fy}/{day}/{run_id}"
    folder = output / relative
    folder.mkdir(parents=True, exist_ok=False)
    prefix = ("Run ID", "Snapshot Date", "Fiscal Year")
    tables = {
        "uncategorized.csv": (DETAIL_FIELDS, report.details),
        "aor_summary.csv": (prefix + AOR_DIMS + METRIC_FIELDS, summarize(report, AOR_DIMS)),
        "component_summary.csv": (prefix + COMPONENT_DIMS + METRIC_FIELDS, summarize(report, COMPONENT_DIMS)),
        "owner_summary.csv": (prefix + OWNER_DIMS + METRIC_FIELDS, summarize(report, OWNER_DIMS)),
        "snapshot.csv": (tuple(report.summary), [report.summary]),
        "rejected_rows.csv": (REJECT_FIELDS, report.rejected),
    }
    for filename, (fields, rows) in tables.items():
        csv_file(folder / filename, fields, rows)
    if excel:
        workbook(folder / "Uncategorized Report.xlsx", report, tables, moderate, critical, sample)
    artifacts = {}
    for path in sorted(folder.iterdir()):
        with path.open("rb") as handle:
            digest = hashlib.file_digest(handle, "sha256").hexdigest()
        artifacts[path.name] = {"bytes": path.stat().st_size, "sha256": digest}
    manifest = {
        "schema_version": 1, "run_id": run_id, "snapshot_date": day,
        "source_data_date": source_date, "fiscal_year": fy, "source_sha256": source_hash,
        "sample_data": sample, "snapshot_folder": relative + "/",
        "moderate_starts_at": moderate, "critical_starts_at": critical,
        "uncategorized_items": len(report.details), "rejected_rows": len(report.rejected),
        "artifacts": artifacts,
    }
    content = json.dumps(manifest, indent=2) + "\n"
    (folder / "manifest.json").write_text(content, encoding="utf-8")
    latest = output / "latest"
    latest.mkdir(parents=True, exist_ok=True)
    # Readers first resolve this pointer, then read immutable files from one run.
    descriptor, temporary = tempfile.mkstemp(prefix="manifest.", dir=latest)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            handle.write(content)
        os.replace(temporary, latest / "manifest.json")
    finally:
        Path(temporary).unlink(missing_ok=True)
    return folder
