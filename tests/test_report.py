import csv
from contextlib import redirect_stderr, redirect_stdout
import hashlib
import io
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from bwc_report.__main__ import main
from bwc_report.core import (REQUIRED, ReportError, age_band, fiscal_year, organization,
                             read_csv, summarize, transform)
from bwc_report.export import publish_local


DAY = date(2026, 10, 13)


def row(evidence_id="0001", uploaded="2026-09-13", **changes):
    return {"Evidence ID": evidence_id, "Categories": "", "Status": "Active",
            "Date Uploaded": uploaded, "Owner Groups": "HSI Baltimore Users, Administrators",
            "Owner Badge ID": "0007", "Owner First Name": "Demo", "Owner Last Name": "Person",
            "Owner Role": "Officer", "__row__": "2", **changes}


class RuleTests(unittest.TestCase):
    def test_selected_boundaries(self):
        expected = {0: "Normal", 14: "Normal", 15: "Moderate", 29: "Moderate",
                    30: "Critical", 31: "Critical"}
        for days, band in expected.items():
            with self.subTest(days=days):
                self.assertEqual(age_band(days)[0], band)

    def test_invalid_thresholds_and_negative_age(self):
        for args in ((1, 30, 15), (-1, 15, 30), (0, 0, 30)):
            with self.assertRaises(ReportError):
                age_band(*args)

    def test_fiscal_year_is_exact_october_rollover(self):
        self.assertEqual(fiscal_year(date(2026, 9, 30)), 2026)
        self.assertEqual(fiscal_year(date(2026, 10, 1)), 2027)
        self.assertEqual(fiscal_year(date(2026, 12, 31)), 2027)

    def test_organization_rules_and_unknowns(self):
        cases = {
            "HSI Baltimore Users, Admin": ("HSI", "HSI Baltimore", "AOR Baltimore"),
            "ERO Miami 2nd Line Sup, Admin": ("ERO", "ERO Miami", "AOR Miami"),
            "FLETC Users": ("FLETC", "FLETC", "FLETC"),
            "FOIA Users": ("FOIA", "FOIA", "FOIA"),
            "Office of Professional Responsibility Users": ("OPR", "Office of Professional Responsibility", "Office of Professional Responsibility"),
            "Office of Firearms and Tactical Programs": ("OFTP", "Office of Firearms and Tactical Programs", "Office of Firearms and Tactical Programs"),
            "": ("Unassigned", "Unassigned", "Unassigned"),
            "Legal Users": ("Other", "Legal", "Legal"),
        }
        for source, expected in cases.items():
            with self.subTest(source=source):
                self.assertEqual(organization(source), expected)

    def test_status_and_category_filters(self):
        report = transform([
            row("1"), row("2", Status="Queued for Deletion"),
            row("3", Status="Deleted"), row("4", Categories="Case evidence"),
        ], DAY, "test")
        self.assertEqual([item["Evidence ID"] for item in report.details], ["1", "2"])
        self.assertEqual(report.summary["Excluded Status Rows"], 1)
        self.assertEqual(report.summary["Categorized Rows"], 1)

    def test_invalid_and_future_dates_are_rejected(self):
        report = transform([row("1", "2026-10-14"), row("2", "not a date"),
                            row("", "2026-10-12")], DAY, "test")
        self.assertEqual(report.summary["Total Uncategorized Items"], 0)
        self.assertEqual(report.summary["Rejected Rows"], 3)
        self.assertEqual(len(report.rejected), 3)

    def test_optional_created_date_does_not_lose_valid_record(self):
        report = transform([row(**{"Created Date (Record Start)": "not a date"})], DAY, "test")
        self.assertEqual(len(report.details), 1)
        self.assertIn("Invalid created date", report.details[0]["Data Quality Notes"])

    def test_duplicate_rows_deduplicate_but_conflicts_block_run(self):
        report = transform([row(), row(**{"__row__": "3"})], DAY, "test")
        self.assertEqual(len(report.details), 1)
        self.assertEqual(report.summary["Duplicate Rows Removed"], 1)
        with self.assertRaises(ReportError):
            transform([row(), row(Categories="categorized")], DAY, "test")

    def test_zero_records_produce_snapshot(self):
        report = transform([], DAY, "zero")
        self.assertEqual(report.details, [])
        self.assertEqual(report.summary["Total Uncategorized Items"], 0)
        self.assertEqual(report.summary["Critical Items"], 0)
        self.assertEqual(summarize(report, ("AOR",)), [])

    def test_unknown_and_blank_groups_remain_in_totals(self):
        report = transform([row("1", **{"Owner Groups": ""}),
                            row("2", **{"Owner Groups": "Legal Users"})], DAY, "test")
        summary = summarize(report, ("Component",))
        self.assertEqual(sum(item["Total Uncategorized Items"] for item in summary), 2)
        self.assertEqual({item["Component"] for item in summary}, {"Other", "Unassigned"})

    def test_conflicting_officer_names_do_not_use_first_match(self):
        report = transform([row("1"), row("2", **{"Owner First Name": "Another Demo"})], DAY, "test")
        summary = summarize(report, ("Owner Badge ID", "Owner First Name", "Component Name"))
        self.assertEqual(len(summary), 2)
        self.assertEqual(sum(item["Total Uncategorized Items"] for item in summary), 2)


class InputOutputTests(unittest.TestCase):
    def csv_bytes(self, rows, headers=REQUIRED):
        text = io.StringIO(newline="")
        writer = csv.DictWriter(text, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
        return text.getvalue().encode("utf-8-sig")

    def test_bom_nulls_blank_row_and_last_real_row(self):
        content = self.csv_bytes([row("0001", Categories="NULL"),
                                  {field: "" for field in REQUIRED}, row("0002")])
        report = transform(read_csv(content), DAY, "test")
        self.assertEqual(len(report.details), 2)
        self.assertEqual(report.summary["Blank Rows Removed"], 1)
        self.assertEqual(report.details[-1]["Evidence ID"], "0002")

    def test_missing_and_duplicate_headers(self):
        with self.assertRaisesRegex(ReportError, "Missing required"):
            read_csv(b"Evidence ID,Status\n1,Active\n")
        with self.assertRaisesRegex(ReportError, "duplicate column"):
            read_csv(b"Evidence ID,Evidence ID\n1,1\n")

    def test_vendor_column_mapping(self):
        content = self.csv_bytes([{"Id": "0001", **row()}], ("Id",) + REQUIRED[1:])
        parsed = read_csv(content, column_map={"Evidence ID": "Id"})
        self.assertEqual(parsed[0]["Evidence ID"], "0001")

    def test_malformed_rows_fail(self):
        content = b"Evidence ID,Categories,Status,Date Uploaded,Owner Groups\n1,,Active,2026-10-13,HSI,extra\n"
        with self.assertRaisesRegex(ReportError, "number of cells"):
            read_csv(content)

    def test_manifest_hashes_and_atomic_latest_pointer(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            for run in ("first", "second"):
                folder = publish_local(transform([row()], DAY, run), output, "source-hash", None)
                manifest = json.loads((folder / "manifest.json").read_text())
                for name, metadata in manifest["artifacts"].items():
                    self.assertEqual(hashlib.sha256((folder / name).read_bytes()).hexdigest(), metadata["sha256"])
            latest = json.loads((output / "latest/manifest.json").read_text())
            self.assertEqual(latest["run_id"], "second")
            self.assertTrue((output / "snapshots/FY2027/2026-10-13/first/uncategorized.csv").exists())

    def test_excel_can_open_and_source_formula_stays_literal(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            report = transform([row(**{"Owner First Name": '=HYPERLINK("https://example.invalid")'})], DAY, "excel")
            folder = publish_local(report, Path(directory), "hash", None, excel=True, sample=True)
            book = load_workbook(folder / "Uncategorized Report.xlsx")
            sheet = book["Uncategorized"]
            headers = [cell.value for cell in sheet[1]]
            cell = sheet.cell(2, headers.index("Owner First Name") + 1)
            self.assertEqual(cell.data_type, "s")
            self.assertEqual(sheet.cell(2, headers.index("Days Since Upload") + 1).fill.fgColor.rgb, "00C41230")
            self.assertEqual(len(book.sheetnames), 7)
            self.assertIn("Uncategorized", sheet.tables)
            book.close()

    def test_empty_excel_contains_zero_snapshot_and_headers(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            folder = publish_local(transform([], DAY, "zero"), Path(directory), "hash", None, excel=True)
            book = load_workbook(folder / "Uncategorized Report.xlsx")
            self.assertEqual(book["Uncategorized"].max_row, 1)
            self.assertEqual(book["Snapshot"].max_row, 2)
            book.close()

    def test_cli_reject_failure_preserves_previous_latest(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "input.csv"
            source.write_bytes(self.csv_bytes([row()]))
            args = ["--input", str(source), "--output", str(root / "output"), "--as-of", DAY.isoformat()]
            with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                self.assertEqual(main(args), 0)
            before = (root / "output/latest/manifest.json").read_bytes()
            source.write_bytes(self.csv_bytes([row(uploaded="2026-10-14")]))
            with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                self.assertEqual(main(args + ["--fail-on-rejects"]), 1)
            self.assertEqual((root / "output/latest/manifest.json").read_bytes(), before)

    def test_shared_synthetic_acceptance_fixture(self):
        root = Path(__file__).resolve().parents[1]
        report = transform(read_csv((root / "examples/evidence.sample.csv").read_bytes()), DAY, "sample")
        expected = json.loads((root / "examples/expected-summary.json").read_text())
        for field, value in expected.items():
            with self.subTest(field=field):
                self.assertEqual(report.summary[field], value)
        self.assertEqual(report.summary["Average Days Since Upload"], 29.125)


if __name__ == "__main__":
    unittest.main()
